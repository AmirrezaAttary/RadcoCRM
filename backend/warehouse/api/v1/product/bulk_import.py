import re
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import IntegrityError, transaction
from openpyxl import load_workbook

from rest_framework import serializers, status
from rest_framework.generics import GenericAPIView
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from access.permissions import WarehousePermission
from warehouse.models.product import Product, ProductCategory, UNIT_TYPE
from warehouse.models.unit_measurement import (
    AreaUnit,
    CountUnit,
    DimensionUnit,
    VolumeUnit,
    WeightUnit,
)
from warehouse.models.warehouse import Warehouse
from warehouse.models.warehouse_process import InventoryStock
from warehouse.models.warehouse_staff import WarehouseStaff
from warehouse.services.warehouse_process import InventoryService


MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 10000
MAX_PRICE_DIGITS = 12  # مطابق Product.sale_price


class ProductExcelImportSerializer(serializers.Serializer):
    file = serializers.FileField()
    # فقط وقتی انبار در آدرس آمده لازم است (ثبت‌کننده‌ی موجودی اول دوره)
    performed_by_id = serializers.IntegerField(required=False)

    def validate_file(self, value):
        if value.size > MAX_FILE_SIZE:
            raise serializers.ValidationError(
                "حجم فایل نباید بیشتر از ۱۰ مگابایت باشد."
            )

        if not value.name.lower().endswith(".xlsx"):
            raise serializers.ValidationError(
                "فقط فایل Excel با پسوند .xlsx پشتیبانی می‌شود."
            )

        return value


# ───────────────────────── ستون‌ها ─────────────────────────

HEADER_ALIASES = {
    "name": {
        "name", "product_name", "product",
        "نام کالا", "نام محصول", "کالا", "نام",
    },
    "sale_price": {
        "sale_price", "price", "قیمت فروش", "قیمت",
    },
    "category": {
        "category", "category_name",
        "دسته بندی", "دسته‌بندی", "دسته", "گروه کالا",
    },
    "unit_type": {
        "unit_type", "unit", "نوع واحد", "واحد",
    },
    # dimension
    "meter": {"meter", "متر"},
    "centimeter": {"centimeter", "سانتیمتر", "سانتی متر"},
    "inch": {"inch", "اینچ"},
    # weight
    "gram": {"gram", "گرم"},
    "kilogram": {"kilogram", "کیلوگرم", "کیلو"},
    "milligram": {"milligram", "میلی گرم", "میلی‌گرم"},
    "ton": {"ton", "تن"},
    # volume
    "liter": {"liter", "litre", "لیتر"},
    "milliliter": {"milliliter", "میلی لیتر", "میلی‌لیتر"},
    "gallon": {"gallon", "گالن"},
    # area
    "square_meter": {"square_meter", "متر مربع"},
    "square_centimeter": {"square_centimeter", "سانتیمتر مربع", "سانتی متر مربع"},
    # count
    "quantity_per_unit": {
        "quantity_per_unit", "quantity", "تعداد در هر بسته", "تعداد در بسته", "تعداد",
    },
    # موجودی انبار (فقط وقتی id انبار در آدرس باشد)
    "initial_quantity": {
        "initial_quantity", "initial_stock", "opening_stock",
        "موجودی اولیه", "موجودی اول دوره", "موجودی افتتاحیه", "موجودی",
    },
    "minimum_stock": {
        "minimum_stock", "min_stock", "min",
        "حداقل موجودی", "حداقل",
    },
    "maximum_stock": {
        "maximum_stock", "max_stock", "max",
        "حداکثر موجودی", "حداکثر",
    },
}

REQUIRED_FIELDS = {"name", "sale_price"}

STOCK_FIELDS = ("initial_quantity", "minimum_stock", "maximum_stock")

REQUIRED_LABELS = {
    "name": "نام کالا",
    "sale_price": "قیمت فروش",
    "initial_quantity": "موجودی اولیه",
    "minimum_stock": "حداقل موجودی",
    "maximum_stock": "حداکثر موجودی",
}

# هر نوع واحد → (مدل، فیلد FK روی Product، فیلدهای عددی)
UNIT_CONFIG = {
    UNIT_TYPE.DIMENSION: (
        DimensionUnit, "dimension_unit", ("meter", "centimeter", "inch"),
    ),
    UNIT_TYPE.WEIGHT: (
        WeightUnit, "weight_unit", ("gram", "kilogram", "milligram", "ton"),
    ),
    UNIT_TYPE.VOLUME: (
        VolumeUnit, "volume_unit", ("liter", "milliliter", "gallon"),
    ),
    UNIT_TYPE.AREA: (
        AreaUnit, "area_unit", ("square_meter", "square_centimeter"),
    ),
    UNIT_TYPE.COUNT: (
        CountUnit, "count_unit", ("quantity_per_unit",),
    ),
}

UNIT_TYPE_ALIASES = {
    UNIT_TYPE.DIMENSION: {"dimension", "dimensions", "ابعاد", "بعد"},
    UNIT_TYPE.WEIGHT: {"weight", "وزن"},
    UNIT_TYPE.VOLUME: {"volume", "حجم"},
    UNIT_TYPE.AREA: {"area", "مساحت"},
    UNIT_TYPE.COUNT: {"count", "تعداد"},
}


# ───────────────────────── توابع کمکی ─────────────────────────

def normalize_header(value):
    value = str(value or "").strip().lower()
    value = value.replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")
    return re.sub(r"[\s\-_]+", " ", value)


def normalize_digits(value):
    text = str(value if value is not None else "").strip()
    table = str.maketrans(
        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
        "01234567890123456789",
    )
    return text.translate(table)


def normalize_text(value):
    text = str(value if value is not None else "").strip()
    text = text.replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"\s+", " ", text)


def parse_price(value):
    """قیمت صحیح (بدون اعشار) برمی‌گرداند؛ در صورت خطا ValueError."""
    if value is None or str(value).strip() == "":
        raise ValueError("قیمت فروش خالی است.")

    if isinstance(value, bool):
        raise ValueError("قیمت فروش معتبر نیست.")

    if isinstance(value, (int, float, Decimal)):
        text = format(Decimal(str(value)), "f")
    else:
        text = normalize_digits(value)
        text = re.sub(r"[,٬،\s]", "", text)
        text = re.sub(r"(تومان|ریال|toman|rial)$", "", text, flags=re.I)

    try:
        number = Decimal(text)
    except InvalidOperation:
        raise ValueError("قیمت فروش باید عدد باشد.")

    if number < 0:
        raise ValueError("قیمت فروش نمی‌تواند منفی باشد.")

    if number != number.to_integral_value():
        raise ValueError("قیمت فروش باید عدد صحیح (بدون اعشار) باشد.")

    number = number.to_integral_value()

    if len(str(int(number))) > MAX_PRICE_DIGITS:
        raise ValueError("قیمت فروش بیش از حد بزرگ است.")

    return number


def parse_unit_type(value):
    if value is None or str(value).strip() == "":
        return None

    text = normalize_text(value).lower()

    for unit_type, aliases in UNIT_TYPE_ALIASES.items():
        if text in aliases:
            return unit_type

    raise ValueError(
        "نوع واحد باید یکی از این‌ها باشد: ابعاد، وزن، حجم، مساحت، تعداد."
    )


FIELD_LABELS = {
    "meter": "متر", "centimeter": "سانتیمتر", "inch": "اینچ",
    "gram": "گرم", "kilogram": "کیلوگرم", "milligram": "میلی‌گرم", "ton": "تن",
    "liter": "لیتر", "milliliter": "میلی‌لیتر", "gallon": "گالن",
    "square_meter": "متر مربع", "square_centimeter": "سانتیمتر مربع",
    "quantity_per_unit": "تعداد در هر بسته",
    "initial_quantity": "موجودی اولیه",
    "minimum_stock": "حداقل موجودی",
    "maximum_stock": "حداکثر موجودی",
}


def parse_unit_number(value, field, integer=False):
    label = FIELD_LABELS.get(field, field)
    if value is None or str(value).strip() == "":
        return None

    if isinstance(value, bool):
        raise ValueError(f"مقدار ستون «{label}» عدد معتبر نیست.")

    if isinstance(value, (int, float)):
        number = Decimal(str(value))
    else:
        text = normalize_digits(value)
        text = re.sub(r"[,٬،\s]", "", text).replace("٫", ".")
        try:
            number = Decimal(text)
        except InvalidOperation:
            raise ValueError(f"مقدار ستون «{label}» عدد معتبر نیست.")

    if not number.is_finite():
        raise ValueError(f"مقدار ستون «{label}» عدد معتبر نیست.")

    if number < 0:
        raise ValueError(f"مقدار ستون «{label}» نمی‌تواند منفی باشد.")

    if integer:
        if number != number.to_integral_value():
            raise ValueError(f"مقدار ستون «{label}» باید عدد صحیح باشد.")
        return int(number)

    return float(number)


def parse_stock_values(item):
    """موجودی اولیه، حداقل و حداکثر را می‌خواند و مثل InitialStockSerializer اعتبارسنجی می‌کند."""
    values = {}
    for field in STOCK_FIELDS:
        raw = item.get(field)
        if raw is None or str(raw).strip() == "":
            raise ValueError(f"مقدار ستون «{FIELD_LABELS[field]}» خالی است.")
        values[field] = parse_unit_number(raw, field)

    quantity = values["initial_quantity"]
    minimum = values["minimum_stock"]
    maximum = values["maximum_stock"]

    if quantity <= 0:
        raise ValueError("موجودی اولیه باید بیشتر از صفر باشد.")
    if maximum < minimum:
        raise ValueError("حداکثر موجودی نمی‌تواند کمتر از حداقل موجودی باشد.")
    if quantity < minimum:
        raise ValueError(
            f"موجودی اولیه نمی‌تواند کمتر از حداقل موجودی ({minimum:g}) باشد."
        )
    if quantity > maximum:
        raise ValueError(
            f"موجودی اولیه نمی‌تواند بیشتر از حداکثر موجودی ({maximum:g}) باشد."
        )

    return values


def cell_text(value):
    """مقدار سلول را به رشته تمیز تبدیل می‌کند (۱۰.۰ → ۱۰)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


# ───────────────────────── View ─────────────────────────

class ProductExcelImportApi(GenericAPIView):
    """
    ایمپورت کالا از اکسل.

    - POST /products/bulk-import/                  → فقط ثبت کالا (مثل قبل)
    - POST /products/bulk-import/<warehouse_id>/   → ثبت کالا + موجودی اول دوره،
      حداقل و حداکثر موجودی در همان انبار (فیلد فرم: performed_by_id هم لازم است)
    """

    serializer_class = ProductExcelImportSerializer
    parser_classes = [MultiPartParser, FormParser]

    def get_permissions(self):
        permissions = [IsAuthenticated()]
        # ثبت موجودی اول دوره مثل InitialStockCreateView نیاز به دسترسی انبار دارد
        if self.kwargs.get("warehouse_id") is not None:
            permissions.append(WarehousePermission())
        return permissions

    def post(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        uploaded = serializer.validated_data["file"]

        warehouse = None
        performed_by = None
        warehouse_id = kwargs.get("warehouse_id")

        if warehouse_id is not None:
            warehouse = Warehouse.objects.filter(pk=warehouse_id).first()
            if warehouse is None:
                return Response(
                    {"detail": "انبار یافت نشد."},
                    status=status.HTTP_404_NOT_FOUND,
                )
            if not warehouse.is_active:
                return Response(
                    {"detail": "این انبار غیرفعال است."},
                    status=status.HTTP_400_BAD_REQUEST,
                )

            performed_by_id = serializer.validated_data.get("performed_by_id")
            if performed_by_id is None:
                return Response(
                    {"performed_by_id": ["این فیلد برای ثبت موجودی اول دوره الزامی است."]},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            performed_by = WarehouseStaff.objects.filter(
                pk=performed_by_id, is_active=True
            ).first()
            if performed_by is None:
                return Response(
                    {"detail": "انباردار یافت نشد یا فعال نیست."},
                    status=status.HTTP_404_NOT_FOUND,
                )

        try:
            workbook = load_workbook(
                BytesIO(uploaded.read()),
                read_only=True,
                data_only=True,
            )

            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)

                if not headers:
                    return Response(
                        {"detail": "فایل خالی است یا ردیف عنوان ستون‌ها وجود ندارد."},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                normalized_aliases = {
                    normalize_header(alias): field
                    for field, aliases in HEADER_ALIASES.items()
                    for alias in aliases
                }

                columns = {}
                for index, header in enumerate(headers):
                    key = normalized_aliases.get(normalize_header(header))
                    if key and key not in columns:
                        columns[key] = index

                required_fields = set(REQUIRED_FIELDS)
                if warehouse is not None:
                    required_fields |= set(STOCK_FIELDS)

                if warehouse is None:
                    stock_cols = [f for f in STOCK_FIELDS if f in columns]
                    if stock_cols:
                        return Response(
                            {
                                "detail": (
                                    "فایل ستون موجودی دارد؛ برای ثبت موجودی باید "
                                    "id انبار در آدرس باشد: "
                                    "products/bulk-import/<warehouse_id>/"
                                ),
                                "stock_columns": [
                                    REQUIRED_LABELS[f] for f in stock_cols
                                ],
                            },
                            status=status.HTTP_400_BAD_REQUEST,
                        )

                missing = [
                    f for f in (
                        "name", "sale_price", *STOCK_FIELDS
                    ) if f in required_fields and f not in columns
                ]
                if missing:
                    return Response(
                        {
                            "detail": "ستون‌های ضروری پیدا نشدند.",
                            "missing_columns": [
                                REQUIRED_LABELS[field] for field in missing
                            ],
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                parsed_rows = []
                errors = []

                for excel_row, row in enumerate(rows, start=2):
                    if excel_row > MAX_ROWS + 1:
                        return Response(
                            {
                                "detail": (
                                    f"حداکثر {MAX_ROWS} ردیف "
                                    "قابل پردازش است."
                                )
                            },
                            status=status.HTTP_400_BAD_REQUEST,
                        )

                    if not any(value not in (None, "") for value in row):
                        continue

                    item = {
                        field: row[index] if index < len(row) else None
                        for field, index in columns.items()
                    }

                    name = cell_text(item.get("name"))
                    if not name:
                        errors.append({
                            "row": excel_row,
                            "name": "",
                            "reason": "نام کالا خالی است.",
                        })
                        continue

                    if len(name) > 200:
                        errors.append({
                            "row": excel_row,
                            "name": name[:50],
                            "reason": "نام کالا بیش از ۲۰۰ کاراکتر است.",
                        })
                        continue

                    try:
                        price = parse_price(item.get("sale_price"))
                    except ValueError as exc:
                        errors.append({
                            "row": excel_row,
                            "name": name,
                            "reason": str(exc),
                        })
                        continue

                    category_name = cell_text(item.get("category"))
                    if len(category_name) > 100:
                        errors.append({
                            "row": excel_row,
                            "name": name,
                            "reason": "نام دسته‌بندی بیش از ۱۰۰ کاراکتر است.",
                        })
                        continue

                    try:
                        unit_type = parse_unit_type(item.get("unit_type"))

                        unit_values = {}
                        if unit_type:
                            _, _, fields = UNIT_CONFIG[unit_type]
                            for field in fields:
                                unit_values[field] = parse_unit_number(
                                    item.get(field),
                                    field,
                                    integer=(field == "quantity_per_unit"),
                                )
                    except ValueError as exc:
                        errors.append({
                            "row": excel_row,
                            "name": name,
                            "reason": str(exc),
                        })
                        continue

                    stock_values = None
                    if warehouse is not None:
                        try:
                            stock_values = parse_stock_values(item)
                        except ValueError as exc:
                            errors.append({
                                "row": excel_row,
                                "name": name,
                                "reason": str(exc),
                            })
                            continue

                    parsed_rows.append({
                        "row": excel_row,
                        "name": name,
                        "sale_price": price,
                        "category_name": category_name,
                        "unit_type": unit_type,
                        "unit_values": unit_values,
                        "stock_values": stock_values,
                    })

            finally:
                workbook.close()

        except Exception:
            return Response(
                {
                    "detail": (
                        "خواندن فایل اکسل ممکن نشد. "
                        "سالم بودن فایل و فرمت .xlsx را بررسی کنید."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ── تشخیص تکراری‌ها (نام + دسته‌بندی) ──────────────────────
        # کالای تکراری = هم‌نام و هم‌دسته.
        # - بدون انبار: کالای موجود در دیتابیس ثبت نمی‌شود (مثل قبل).
        # - با انبار: کالای موجود دوباره ساخته نمی‌شود و فقط موجودی همین انبار
        #   برایش ثبت می‌شود (برای کالای مشترک بین چند انبار)؛ اگر موجودی
        #   آن در این انبار قبلاً ثبت شده باشد، ردیف تکراری حساب می‌شود.
        def dup_key(name, category_name):
            return (name.lower(), category_name.lower())

        existing_categories = {
            normalize_text(c.name).lower(): c
            for c in ProductCategory.objects.all()
        }

        existing_products = {}
        names = {r["name"] for r in parsed_rows}
        for product in Product.objects.filter(name__in=names).select_related("category"):
            key = dup_key(
                normalize_text(product.name),
                normalize_text(product.category.name if product.category else ""),
            )
            existing_products.setdefault(key, product)

        stocked_product_ids = set()
        if warehouse is not None:
            stocked_product_ids = set(
                InventoryStock.objects
                .filter(warehouse=warehouse)
                .values_list("product_id", flat=True)
            )

        seen = set()
        to_create = []
        skipped_duplicates = []

        for r in parsed_rows:
            key = dup_key(r["name"], r["category_name"])

            if key in seen:
                skipped_duplicates.append({
                    "row": r["row"],
                    "name": r["name"],
                    "reason": "کالا با همین نام و دسته‌بندی در همین فایل تکراری است؛ ثبت نشد.",
                })
                continue
            seen.add(key)

            product = existing_products.get(key)
            r["product"] = product

            if product is not None:
                if warehouse is None:
                    skipped_duplicates.append({
                        "row": r["row"],
                        "name": r["name"],
                        "reason": "کالا با همین نام و دسته‌بندی تکراری است؛ ثبت نشد.",
                    })
                    continue
                if product.pk in stocked_product_ids:
                    skipped_duplicates.append({
                        "row": r["row"],
                        "name": r["name"],
                        "reason": "موجودی اول دوره این کالا در این انبار قبلاً ثبت شده است؛ ثبت نشد.",
                    })
                    continue

            to_create.append(r)

        created_categories = []
        stock_items = []
        new_product_count = 0

        try:
            with transaction.atomic():
                for r in to_create:
                    product = r["product"]

                    if product is None:
                        category = None
                        if r["category_name"]:
                            cat_key = r["category_name"].lower()
                            category = existing_categories.get(cat_key)
                            if category is None:
                                category = ProductCategory.objects.create(
                                    name=r["category_name"]
                                )
                                existing_categories[cat_key] = category
                                created_categories.append(category.name)

                        product_kwargs = {
                            "name": r["name"],
                            "sale_price": r["sale_price"],
                            "category": category,
                            "unit_type": r["unit_type"],
                        }

                        if r["unit_type"]:
                            UnitModel, fk_field, _ = UNIT_CONFIG[r["unit_type"]]
                            product_kwargs[fk_field] = UnitModel.objects.create(
                                **r["unit_values"]
                            )

                        product = Product.objects.create(**product_kwargs)
                        r["product_created"] = True
                        new_product_count += 1
                    else:
                        r["product_created"] = False

                    if warehouse is not None:
                        sv = r["stock_values"]
                        stock_items.append({
                            "product": product,
                            "quantity": sv["initial_quantity"],
                            "minimum_stock": sv["minimum_stock"],
                            "maximum_stock": sv["maximum_stock"],
                        })

                if warehouse is not None and stock_items:
                    InventoryService.bulk_add_initial_stock(
                        warehouse=warehouse,
                        items=stock_items,
                        performed_by=performed_by,
                    )
        except (ValueError, IntegrityError) as exc:
            # هیچ‌چیز ثبت نشده است (کل ایمپورت rollback شد)
            detail = str(exc) if isinstance(exc, ValueError) else (
                "ثبت موجودی با داده‌ی موجود تداخل دارد (احتمالاً هم‌زمان ثبت شده)."
            )
            return Response(
                {"detail": detail},
                status=status.HTTP_409_CONFLICT,
            )

        response_data = {
            "detail": "پردازش فایل اکسل انجام شد.",
            "created_count": new_product_count,
            "duplicate_count": len(skipped_duplicates),
            "error_count": len(errors),
            "created_categories": created_categories,
            "created": [],
            "duplicates": skipped_duplicates,
            "errors": errors,
        }

        for r in to_create:
            entry = {
                "name": r["name"],
                "sale_price": int(r["sale_price"]),
                "category": r["category_name"],
            }
            if warehouse is not None:
                entry.update({
                    "product_created": r["product_created"],
                    "initial_quantity": r["stock_values"]["initial_quantity"],
                    "minimum_stock": r["stock_values"]["minimum_stock"],
                    "maximum_stock": r["stock_values"]["maximum_stock"],
                })
            response_data["created"].append(entry)

        if warehouse is not None:
            response_data["warehouse"] = {
                "id": warehouse.pk,
                "name": warehouse.name,
                "code": warehouse.code,
            }
            response_data["stock_registered_count"] = len(stock_items)
            response_data["existing_product_count"] = (
                len(stock_items) - new_product_count
            )

        return Response(response_data, status=status.HTTP_200_OK)
