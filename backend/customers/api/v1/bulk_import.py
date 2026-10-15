
import re
from io import BytesIO

from django.db import transaction
from openpyxl import load_workbook

from rest_framework import status
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.generics import GenericAPIView

from customers.models import Customer
from .serializers import CustomerExcelImportSerializer


MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 10000

HEADER_ALIASES = {
    "full_name": {
        "full_name", "fullname", "name", "customer_name",
        "نام", "نام و نام خانوادگی", "نام مشتری", "نام کامل",
    },
    "phone_number": {
        "phone_number", "phone", "mobile", "mobile_number",
        "telephone", "شماره تماس", "شماره تلفن",
        "موبایل", "شماره موبایل", "تلفن",
    },
    "job_title": {"job_title", "job", "position", "سمت شغلی", "سمت"},
    "company_name": {
        "company_name", "company", "organization",
        "نام شرکت", "شرکت", "سازمان",
    },
    "address": {"address", "آدرس", "نشانی"},
    "source": {"source", "lead_source", "منبع آشنایی", "منبع"},
    "description": {"description", "notes", "توضیحات", "شرح"},
    "status": {"status", "وضعیت", "وضعیت مشتری"},
}

REQUIRED_FIELDS = {"full_name", "phone_number"}


def normalize_header(value):
    value = str(value or "").strip().lower()
    value = value.replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"[\s\-_]+", " ", value)


def normalize_digits(value):
    text = str(value or "").strip()
    table = str.maketrans(
        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
        "01234567890123456789",
    )
    return text.translate(table)


def normalize_phone(value):
    phone = normalize_digits(value)

    if re.fullmatch(r"\d+\.0", phone):
        phone = phone[:-2]

    phone = re.sub(r"[\s()\-+]", "", phone)

    if phone.startswith("0098"):
        phone = "0" + phone[4:]
    elif phone.startswith("98") and len(phone) >= 12:
        phone = "0" + phone[2:]

    return phone


def parse_status(value):
    if value is None or str(value).strip() == "":
        return Customer.CustomerStatus.POTENTIAL

    text = normalize_digits(value).strip().lower()
    text = text.replace("ي", "ی").replace("ك", "ک")

    if text in {"1", "بالقوه", "بالقوه", "مشتری بالقوه", "potential"}:
        return Customer.CustomerStatus.POTENTIAL

    if text in {"2", "بالفعل", "مشتری بالفعل", "active"}:
        return Customer.CustomerStatus.ACTIVE

    raise ValueError(
        "وضعیت باید ۱ (مشتری بالقوه) یا ۲ (مشتری بالفعل) باشد."
    )


class CustomerExcelImportApi(GenericAPIView):
    serializer_class = CustomerExcelImportSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        uploaded = serializer.validated_data["file"]

        try:
            workbook = load_workbook(
                BytesIO(uploaded.read()),
                read_only=True,
                data_only=True,
            )

            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)

                if not headers:
                    return Response(
                        {"detail": "فایل خالی است یا ردیف عنوان ستون‌ها وجود ندارد."},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                normalized_aliases = {
                    normalize_header(alias): field
                    for field, aliases in HEADER_ALIASES.items()
                    for alias in aliases
                }

                columns = {}
                for index, header in enumerate(headers):
                    key = normalized_aliases.get(normalize_header(header))
                    if key and key not in columns:
                        columns[key] = index

                missing = sorted(REQUIRED_FIELDS - set(columns))

                if missing:
                    labels = {
                        "full_name": "نام و نام خانوادگی",
                        "phone_number": "شماره تماس",
                    }
                    return Response(
                        {
                            "detail": "ستون‌های ضروری پیدا نشدند.",
                            "missing_columns": [
                                labels[field] for field in missing
                            ],
                        },
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                imported_rows = []
                errors = []

                for excel_row, row in enumerate(rows, start=2):
                    if excel_row > MAX_ROWS + 1:
                        return Response(
                            {
                                "detail": (
                                    f"حداکثر {MAX_ROWS} ردیف "
                                    "قابل پردازش است."
                                )
                            },
                            status=status.HTTP_400_BAD_REQUEST,
                        )

                    if not any(value not in (None, "") for value in row):
                        continue

                    item = {
                        field: row[index] if index < len(row) else None
                        for field, index in columns.items()
                    }

                    full_name = str(item.get("full_name") or "").strip()
                    phone = normalize_phone(item.get("phone_number"))

                    if not full_name:
                        errors.append({
                            "row": excel_row,
                            "phone_number": phone,
                            "reason": "نام مشتری خالی است.",
                        })
                        continue

                    if not phone:
                        errors.append({
                            "row": excel_row,
                            "phone_number": "",
                            "reason": "شماره تماس خالی است.",
                        })
                        continue

                    if not re.fullmatch(r"\d{10,15}", phone):
                        errors.append({
                            "row": excel_row,
                            "phone_number": phone,
                            "reason": "فرمت شماره تماس معتبر نیست.",
                        })
                        continue

                    customer = {
                        "full_name": full_name,
                        "phone_number": phone,
                        "job_title": str(
                            item.get("job_title") or ""
                        ).strip(),
                        "company_name": str(
                            item.get("company_name") or ""
                        ).strip(),
                        "address": str(
                            item.get("address") or ""
                        ).strip(),
                        "source": str(
                            item.get("source") or ""
                        ).strip(),
                        "description": str(
                            item.get("description") or ""
                        ).strip(),
                        "status": Customer.CustomerStatus.POTENTIAL,
                    }

                    try:
                        customer["status"] = parse_status(
                            item.get("status")
                        )
                    except ValueError as exc:
                        errors.append({
                            "row": excel_row,
                            "phone_number": phone,
                            "reason": str(exc),
                        })
                        continue

                    imported_rows.append({
                        "row": excel_row,
                        "data": customer,
                    })

            finally:
                workbook.close()

        except Exception:
            return Response(
                {
                    "detail": (
                        "خواندن فایل اکسل ممکن نشد. "
                        "سالم بودن فایل و فرمت .xlsx را بررسی کنید."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        phones = [
            item["data"]["phone_number"]
            for item in imported_rows
        ]

        existing = set(
            Customer.objects.filter(
                phone_number__in=phones
            ).values_list("phone_number", flat=True)
        )

        seen = set(existing)
        to_create = []
        skipped_duplicates = []

        for item in imported_rows:
            customer = item["data"]
            phone = customer["phone_number"]

            if phone in seen:
                skipped_duplicates.append({
                    "row": item["row"],
                    "phone_number": phone,
                    "reason": "شماره تماس تکراری است؛ این مشتری ثبت نشد.",
                })
                continue

            seen.add(phone)
            to_create.append(
                Customer(created_by=request.user, **customer)
            )

        with transaction.atomic():
            Customer.objects.bulk_create(
                to_create,
                batch_size=500,
            )

        return Response(
            {
                "detail": "پردازش فایل اکسل انجام شد.",
                "created_count": len(to_create),
                "duplicate_count": len(skipped_duplicates),
                "error_count": len(errors),
                "created": [
                    {
                        "full_name": obj.full_name,
                        "phone_number": obj.phone_number,
                    }
                    for obj in to_create
                ],
                "duplicates": skipped_duplicates,
                "errors": errors,
            },
            status=status.HTTP_200_OK,
        )
